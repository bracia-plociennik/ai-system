"""Synthetic tracker lifecycle, evidence and recovery; never uses owner runtime."""
from datetime import datetime, timedelta, timezone
import json
from pathlib import Path
import sys
import tempfile
import unittest
import uuid
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from silverbullet_tracker import TaskTracker, TrackerError, split_interval
from task_identity import initialize_markdown
from work_evidence import EvidenceError, canonical_bytes
from work_evidence_ledger import Ledger


class TrackerFixture:
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.w = Path(self.temp.name) / 'sample-workspace'
        self.vault = self.w / 'tasks/silverbullet';(self.vault/'tasks/Client/Project').mkdir(parents=True)
        (self.w/'tasks/task-store.local.md').write_text('---\nschema_version: 1\nenabled: true\nprofile: silverbullet-v1\nvault_path: '+str(self.vault)+'\n---\n')
        home=self.w/'work-evidence';home.mkdir()
        (home/'config.json').write_bytes(canonical_bytes({'schema_version':1,'workspace_id':self.w.name,'ledger_roots':[str(self.w)],'source_roots':{'tracker':str(home),'silverbullet':str(self.vault)},'timezone':'Europe/Warsaw'}))
        (home/'entities.json').write_bytes(canonical_bytes({'schema_version':1,'entities':[],'task_bindings':[]}))
        self.task_id=str(uuid.uuid4());self.other_id=str(uuid.uuid4());self.device=str(uuid.uuid4())
        self.task=self.vault/'tasks/Client/Project/Task.md'
        self.task.write_text('---\ntype: task\ntitle: Synthetic task\ntask_id: '+self.task_id+'\nstatus: open\ncategory: do\ncustom: keep\n---\nNotes stay unchanged.\n')
        self.task.with_name('Other.md').write_text(self.task.read_text().replace(self.task_id,self.other_id))
        self.time=datetime(2026,10,30,10,tzinfo=timezone.utc)
        self.tracker=TaskTracker(self.w,clock=lambda:self.time,evidence_level='local-test')
        self.run=str(uuid.uuid4())

    def event(self, action, run=None, task=None, replay=False, version=None, **extra):
        data={'operation_id':str(uuid.uuid4()),'task_id':task or self.task_id,'run_id':run or self.run,
              'segment_id':str(uuid.uuid4()),'device_id':self.device,'action':action,
              'occurred_at':self.time.isoformat(),'expected_version':self.tracker.state()['version'] if version is None else version,**extra}
        # Segment identifiers belong only to opening actions.
        if action not in {'start','resume'}:data.pop('segment_id')
        return self.tracker.event(data,replay=replay),data

    def stopped(self, replay=False):
        self.event('start',replay=replay);self.time+=timedelta(minutes=10);self.event('stop',replay=replay)

    def preview(self, **extra):
        return self.tracker.preview({'run_id':self.run,'outcome':'Implemented and verified result','client_id':None,'project_id':None,'unassigned_confirmed':True,**extra},'owner')


class TrackerTests(TrackerFixture, unittest.TestCase):
    def test_pause_excludes_breaks_and_source_task_unchanged(self):
        before=self.task.read_bytes();self.event('start');self.time+=timedelta(minutes=10);self.event('pause')
        self.time+=timedelta(minutes=20);self.event('resume');self.time+=timedelta(minutes=5);self.event('stop')
        self.assertEqual(self.preview()['total_seconds'],900);self.assertEqual(self.task.read_bytes(),before)

    def test_switch_requires_manual_pause(self):
        self.event('start');other=str(uuid.uuid4())
        with self.assertRaisesRegex(TrackerError,'pause-current'):self.event('start',other,self.other_id)
        self.event('pause');self.event('start',other,self.other_id)
        self.assertEqual(next(r for r in self.tracker.state()['runs'] if r['state']=='active')['id'],other)

    def test_idempotent_event_and_content_conflict(self):
        first,event=self.event('start');self.assertEqual(self.tracker.event(event,replay=True),first)
        with self.assertRaisesRegex(TrackerError,'content-conflict'):self.tracker.event({**event,'action':'stop'})
        self.assertEqual(len(self.tracker.state()['runs']),1)

    def test_offline_conflict_preserves_both_and_requires_review(self):
        self.event('start');other=str(uuid.uuid4());result,_=self.event('start',other,self.other_id,replay=True,version=0)
        self.assertEqual(result['status'],'conflict');self.assertEqual(len(self.tracker.state()['runs']),2)
        self.event('stop');self.time+=timedelta(seconds=1)
        with self.assertRaisesRegex(TrackerError,'resolve-time-conflicts'):self.preview()
        for run,task in [(self.run,self.task_id),(other,self.other_id)]:
            self.event('resolve',run,task,correction={'2026-10-30':60},reason='Reviewed two offline devices')
        self.assertEqual(self.preview(time_confirmed=True)['total_seconds'],60)

    def test_offline_time_requires_confirmation(self):
        self.stopped(replay=True)
        with self.assertRaisesRegex(TrackerError,'owner-time'):self.preview()
        self.assertEqual(self.preview(time_confirmed=True)['pieces'][0]['basis'],'owner-confirmed')

    def test_restart_preserves_and_flags_active_run(self):
        self.event('start');again=TaskTracker(self.w,clock=lambda:self.time,evidence_level='local-test');again.mark_restarted()
        self.assertEqual(again.state()['runs'][0]['state'],'active');self.assertTrue(again.state()['runs'][0]['review'])

    def test_complete_defaults_done_reopen_never_restarts(self):
        self.event('start');self.time+=timedelta(seconds=60);self.event('complete');self.event('reopen')
        self.assertEqual(self.preview()['status'],'done');self.assertEqual(self.tracker.state()['runs'][0]['state'],'stopped')

    def test_approval_is_separate_exactly_once_and_xlsx_exports(self):
        self.stopped();preview=self.preview();self.assertEqual(len(Ledger(self.w).committed()[0]),0)
        result=self.tracker.approve(preview['preview_id'],preview['preview_hash'],'owner')
        self.assertEqual(result['status'],'committed');self.assertEqual(result['xlsx_status'],'current')
        self.assertEqual(self.tracker.approve(preview['preview_id'],preview['preview_hash'],'owner'),result)
        self.assertEqual(len(Ledger(self.w).committed()[0]),1)

    def test_source_change_and_wrong_principal_block(self):
        self.stopped();preview=self.preview()
        with self.assertRaisesRegex(TrackerError,'owner-conflict'):self.tracker.approve(preview['preview_id'],preview['preview_hash'],'other')
        self.task.write_text(self.task.read_text()+'New note\n')
        with self.assertRaisesRegex(TrackerError,'task-changed'):self.tracker.approve(preview['preview_id'],preview['preview_hash'],'owner')

    def test_zero_correction_keeps_original_measurement_and_reason(self):
        self.stopped()
        with self.assertRaisesRegex(TrackerError,'reason'):self.preview(correction={'2026-10-30':0},time_confirmed=True)
        preview=self.preview(correction={'2026-10-30':0},reason='The timer was left running',time_confirmed=True)
        self.assertEqual(preview['total_seconds'],0);self.assertEqual(preview['pieces'][0]['basis'],'owner-confirmed')
        self.assertEqual(len(self.tracker.state()['runs'][0]['segments']),1)

    def test_rename_retains_work_identity_and_duplicate_id_blocks(self):
        self.stopped();first=self.preview();self.task.rename(self.task.with_name('Renamed.md'))
        second=self.preview();self.assertEqual(first['total_seconds'],second['total_seconds'])
        self.task.write_text(self.task.with_name('Renamed.md').read_text())
        with self.assertRaisesRegex(TrackerError,'ambiguous'):self.preview()

    def test_native_crash_recovery_no_duplicate(self):
        self.stopped();preview=self.preview()
        with self.assertRaises(EvidenceError):self.tracker.approve(preview['preview_id'],preview['preview_hash'],'owner',failpoint='after-manifest')
        self.tracker.approve(preview['preview_id'],preview['preview_hash'],'owner');self.assertEqual(len(Ledger(self.w).committed()[0]),1)

    def test_workbook_failure_preserves_committed_entry(self):
        self.stopped();preview=self.preview()
        with patch('silverbullet_tracker.Workbook.export',side_effect=OSError('closed')):
            result=self.tracker.approve(preview['preview_id'],preview['preview_hash'],'owner')
        self.assertEqual(result['xlsx_status'],'pending');self.assertEqual(len(Ledger(self.w).committed()[0]),1)
        self.assertEqual(self.tracker.refresh_workbook(preview['preview_id'])['xlsx_status'],'current')

    def test_midnight_and_dst_use_utc_elapsed_seconds(self):
        parts=split_interval('2026-10-30T22:30:00+00:00','2026-10-31T00:30:00+00:00')
        self.assertEqual([p['seconds'] for p in parts],[1800,5400])
        self.assertEqual(sum(p['seconds'] for p in split_interval('2026-10-25T02:30:00+02:00','2026-10-25T02:30:00+01:00')),3600)

    def test_legacy_identity_does_not_invent_date_or_touch_body(self):
        before='---\r\ntype: task\r\ncustom: keep\r\n---\r\n## Notes\r\nBody stays.\r\n'
        after=initialize_markdown(before);self.assertNotIn('created_at',after);self.assertTrue(after.endswith('## Notes\r\nBody stays.\r\n'))
        self.assertEqual(initialize_markdown(after),after)
        created=initialize_markdown(before,created=True,at=self.time);self.assertIn('created_at:',created)

    def test_multiple_runs_keep_task_binding_and_unique_native_work_ids(self):
        self.stopped();first=self.preview();self.tracker.approve(first['preview_id'],first['preview_hash'],'owner')
        self.time+=timedelta(minutes=30);self.run=str(uuid.uuid4());self.stopped();second=self.preview()
        self.tracker.approve(second['preview_id'],second['preview_hash'],'owner')
        entries=Ledger(self.w).committed()[0]
        self.assertEqual(len(entries),2);self.assertEqual(len({e['work_id'] for e in entries.values()}),2)
        registry=json.loads((self.w/'work-evidence/entities.json').read_text())
        self.assertEqual(len([b for b in registry['task_bindings'] if b['task_ref']=='silverbullet:'+self.task_id]),1)

    def test_midnight_creates_two_native_days_and_no_pause_time(self):
        self.time=datetime(2026,10,30,22,30,tzinfo=timezone.utc);self.event('start')
        self.time+=timedelta(hours=2);self.event('stop');preview=self.preview()
        self.assertEqual([p['seconds'] for p in preview['pieces']],[1800,5400])
        self.tracker.approve(preview['preview_id'],preview['preview_hash'],'owner')
        self.assertEqual({e['performed_on'] for e in Ledger(self.w).committed()[0].values()},{'2026-10-30','2026-10-31'})

    def test_stale_preview_version_and_hash_block(self):
        self.stopped();preview=self.preview()
        with self.assertRaisesRegex(TrackerError,'stale-preview'):self.tracker.approve(preview['preview_id'],'bad','owner')
        self.event('reopen',run=None)
        with self.assertRaisesRegex(TrackerError,'state-changed'):self.tracker.approve(preview['preview_id'],preview['preview_hash'],'owner')

    def test_orphan_offline_event_is_preserved_for_clock_review(self):
        result,event=self.event('stop',replay=True,version=0)
        self.assertEqual(result['status'],'conflict');run=self.tracker.state()['runs'][0]
        self.assertTrue(run['review']);self.assertEqual(run['events'][0],event)

    def test_backward_clock_and_future_replay_require_review(self):
        self.event('start',replay=True);self.time-=timedelta(minutes=10);self.event('stop',replay=True)
        self.assertTrue(self.tracker.state()['runs'][0]['review'])
        self.assertEqual(self.tracker.correction({'2026-10-25':90000},'DST full day')['days']['2026-10-25'],90000)
        with self.assertRaisesRegex(TrackerError,'invalid-correction'):self.tracker.correction({'2026-03-29':86400},'DST short day')

    def test_workbook_manual_inputs_survive_timer_export(self):
        from openpyxl import load_workbook
        from work_evidence_xlsx import Tracker as Workbook
        book=Workbook(Ledger(self.w));book.export();path=self.w/'work-evidence/time-tracker.xlsx'
        wb=load_workbook(path);wb['Manual Inputs']['A2']='preserve-owner-input';wb.save(path)
        self.stopped();preview=self.preview();self.tracker.approve(preview['preview_id'],preview['preview_hash'],'owner')
        self.assertEqual(load_workbook(path)['Manual Inputs']['A2'].value,'preserve-owner-input')


if __name__=='__main__':unittest.main()
