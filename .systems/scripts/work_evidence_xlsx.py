"""Bounded XLSX view and hash-bound manual intake; never executes workbook code."""
from __future__ import annotations

import copy
import hashlib
import io
import json
import os
import uuid
import zipfile
from pathlib import PurePosixPath

from work_evidence import (ACTUAL_BASES, EvidenceError, actual_seconds, canonical_bytes,
                           digest, file_hash, read_json, require, sha, uuid_id)
from work_evidence_ledger import atomic_write, coverage, now

SHEETS = ("Summary", "Work Entries", "Sessions", "Manual Inputs", "Conflicts",
          "Import Receipts", "Metadata")
MANUAL = ("row_id", "action", "entry_id", "base_revision", "performed_on", "start",
          "end", "timezone", "duration_seconds", "basis", "client_id", "project_id",
          "title_internal", "outcome", "status", "activity_kind", "reason")
MAX_BYTES = 2 * 1024 * 1024


def dependencies():
    import importlib.metadata as metadata
    expected = {"openpyxl": "3.1.5", "defusedxml": "0.7.1", "et-xmlfile": "2.0.0"}
    require(all(metadata.version(k) == v for k, v in expected.items()),
            "unsupported-xlsx-dependency", "runtime")
    from openpyxl.xml import DEFUSEDXML
    require(DEFUSEDXML, "xml-protection-required", "runtime")


def load_book(data):
    dependencies()
    from openpyxl import load_workbook
    require(len(data) <= MAX_BYTES, "workbook-too-large", "workbook")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names, total = set(), 0
            require(len(archive.infolist()) <= 200, "zip-budget", "archive")
            for part in archive.infolist():
                name = part.filename
                require(name not in names and not name.startswith("/") and
                        ".." not in PurePosixPath(name).parts and "\\" not in name,
                        "unsafe-archive", "archive")
                names.add(name)
                total += part.file_size
                require(not part.flag_bits & 1 and part.file_size <= 4 * MAX_BYTES and
                        total <= 8 * MAX_BYTES and part.file_size <= max(1, part.compress_size) * 200,
                        "zip-budget", "archive")
                require(not any(x in name.lower() for x in ("vbaproject", "externallink", "activex", "embeddings")),
                        "active-workbook-content", "archive")
                if name.endswith((".xml", ".rels")):
                    from defusedxml.ElementTree import fromstring
                    xml = fromstring(archive.read(part))
                    for element in xml.iter():
                        require(element.attrib.get("TargetMode") != "External",
                                "external-link", "workbook")
                        require(not any("macroenabled" in str(v).lower() for v in element.attrib.values()),
                                "active-workbook-content", "archive")
            require("[Content_Types].xml" in names, "invalid-workbook", "archive")
        book = load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
        require(tuple(book.sheetnames) == SHEETS, "workbook-schema-conflict", "sheets")
        for sheet in book:
            require(sheet.max_row <= 10000 and sheet.max_column <= 24,
                    "worksheet-budget", "worksheet")
            for row in sheet:
                for cell in row:
                    require(cell.data_type != "f" and cell.hyperlink is None,
                            "active-cell", "cell")
                    require(not isinstance(cell.value, str) or len(cell.value) <= 10000,
                            "cell-budget", "cell")
        return book
    except EvidenceError:
        raise
    except Exception:
        raise EvidenceError("invalid-workbook", "workbook") from None


def sheet_values(sheet):
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    while rows and all(x is None for x in rows[-1]):
        rows.pop()
    return rows


def protected_hash(book):
    return digest({sheet.title: sheet_values(sheet) for sheet in book
                   if sheet.title != "Manual Inputs"})


def literal(sheet, row):
    sheet.append(row)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"
            cell.quotePrefix = cell.value.startswith(("=", "+", "-", "@"))


def make_book(ledger, workbook_id, export_id, manual_rows=(), receipts=()):
    dependencies()
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.worksheet.datavalidation import DataValidation
    entries, sessions, commits = ledger.committed()
    book = Workbook()
    book.remove(book.active)
    for title in SHEETS:
        book.create_sheet(title)
    totals = coverage(list(entries.values()))
    for row in [("Metric", "Value"), ("Workspace", ledger.config["workspace_id"]),
                ("Actual human minutes", totals["actual_human_seconds"] / 60),
                ("Estimated minutes", totals["estimated_seconds"] / 60),
                ("Unknown entries", totals["unknown_entries"]),
                ("Coverage complete", totals["coverage_complete"]),
                ("Coverage", "Committed local ledger only")]:
        literal(book["Summary"], row)
    literal(book["Work Entries"], ("Date", "Owner", "Client", "Project", "Domain scope", "Outcome", "Status",
                                  "Actual minutes", "Estimated minutes", "Time basis", "Entry ID", "Revision"))
    for entry in sorted(entries.values(), key=lambda e: (e["performed_on"], e["entry_id"])):
        if entry.get("voided", False):
            continue
        actual = actual_seconds(entry)
        literal(book["Work Entries"], (entry["performed_on"], entry["owner_workspace"],
                entry["client_id"], entry["project_id"], "Not recorded", entry["outcome"], entry["status"],
                actual / 60 if actual is not None else None,
                entry["duration"]["seconds"] / 60 if entry["duration"]["basis"] == "estimated" else None,
                entry["duration"]["basis"], entry["entry_id"], entry["revision"]))
    literal(book["Sessions"], ("Date", "Session ID", "Actual minutes", "Allocated minutes"))
    for session in sessions.values():
        literal(book["Sessions"], (session["performed_on"], session["session_id"],
                session["duration"]["seconds"] / 60, sum(a["seconds"] for a in session["allocations"]) / 60))
    literal(book["Manual Inputs"], MANUAL)
    for row in manual_rows:
        literal(book["Manual Inputs"], row)
    literal(book["Conflicts"], ("Category", "Row ID", "Required resolution"))
    literal(book["Import Receipts"], ("Transaction ID", "Row ID", "Entry ID", "Revision", "Row hash"))
    for receipt in receipts:
        literal(book["Import Receipts"], receipt)
    for row in [("Field", "Value"), ("schema_version", 1), ("workbook_id", workbook_id),
                ("export_id", export_id), ("ledger_hash", digest(commits)), ("recorded_at", now())]:
        literal(book["Metadata"], row)
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 34
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_options.horizontalCentered = False
        for cell in sheet[1]:
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="245B4C")
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row > 1:
                    cell.font = Font(name="Calibri", size=11, color="164A80" if sheet.title == "Manual Inputs" else "222222")
                    cell.protection = Protection(locked=sheet.title != "Manual Inputs")
                    if cell.row % 2 == 0:
                        cell.fill = PatternFill("solid", fgColor="F1F4F3")
                sheet.column_dimensions[cell.column_letter].width = max(
                    sheet.column_dimensions[cell.column_letter].width or 13,
                    min(44, max(18, len(str(cell.value or "")) + 2)))
            if row[0].row > 1:
                sheet.row_dimensions[row[0].row].height = 42
        for cell in sheet[1]:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column, values in (("B", '"new,correct,void"'), ("J", '"owner-confirmed,measured,estimated,unknown"'),
                           ("O", '"done,partial,blocked,waiting,unknown"')):
        validation = DataValidation(type="list", formula1=values)
        book["Manual Inputs"].add_data_validation(validation)
        validation.add(column + "2:" + column + "10000")
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


class Tracker:
    def __init__(self, ledger):
        self.ledger = ledger
        self.path = ledger.runtime("time-tracker.xlsx")

    def closed(self):
        require(not self.path.with_name("~$" + self.path.name).exists(), "workbook-open", "workbook")
        # Office's owner file is a heuristic. Caller must also close other editors.

    def snapshot(self, *, allow_stale_ledger=False):
        self.closed()
        data = self.path.read_bytes()
        book = load_book(data)
        meta = dict(sheet_values(book["Metadata"])[1:])
        uuid_id(meta.get("workbook_id"), "workbook_id")
        uuid_id(meta.get("export_id"), "export_id")
        descriptor = read_json(self.ledger.runtime("xlsx/exports/" + meta["export_id"] + ".json", exists=True))
        require(descriptor["workbook_id"] == meta["workbook_id"] and
                descriptor["protected_hash"] == protected_hash(book),
                "generated-sheet-conflict", "workbook")
        require(allow_stale_ledger or meta["ledger_hash"] == self.ledger.snapshot()["ledger_hash"],
                "stale-workbook-ledger", "workbook")
        require(sheet_values(book["Manual Inputs"])[0] == list(MANUAL),
                "manual-header-conflict", "headers")
        expected = hashlib.sha256(data).hexdigest()
        require(file_hash(self.path) == expected, "stale-workbook", "workbook")
        return book, meta, expected

    def write_view(self, workbook_id, manual_rows=(), receipts=(), *, expected=None):
        self.closed()
        require((not self.path.exists() and expected is None) or
                (self.path.exists() and file_hash(self.path) == expected),
                "stale-workbook", "workbook")
        export_id = str(uuid.uuid4())
        data = make_book(self.ledger, workbook_id, export_id, manual_rows, receipts)
        book = load_book(data)
        atomic_write(self.ledger.runtime("xlsx/backups/" + hashlib.sha256(data).hexdigest() + ".xlsx"),
                     data, immutable=True)
        descriptor = {"workbook_id": workbook_id, "protected_hash": protected_hash(book)}
        atomic_write(self.ledger.runtime("xlsx/exports/" + export_id + ".json"),
                     canonical_bytes(descriptor), immutable=True)
        # Recheck after rendering; an open editor/new save must not be overwritten.
        self.closed()
        if expected is not None:
            require(file_hash(self.path) == expected, "stale-workbook", "workbook")
            atomic_write(self.ledger.runtime("xlsx/backups/" + expected + ".xlsx"),
                         self.path.read_bytes(), immutable=True)
        atomic_write(self.path, data)
        return {"status": "exported", "workbook_id": workbook_id, "sha256": file_hash(self.path)}

    def export(self):
        with self.ledger.lock():
            if self.path.exists():
                book, meta, expected = self.snapshot(allow_stale_ledger=True)
                receipts = sheet_values(book["Import Receipts"])[1:]
                acknowledged_transactions = {row[0] for row in receipts}
                _, _, commits = self.ledger.committed()
                for commit in commits:
                    pending = self.ledger.runtime("xlsx/pending/" + commit["transaction_id"] + ".json")
                    require(not pending.exists() or commit["transaction_id"] in acknowledged_transactions,
                            "xlsx-recovery-required", "transaction")
                return self.write_view(meta["workbook_id"], sheet_values(book["Manual Inputs"])[1:],
                                       receipts, expected=expected)
            return self.write_view(str(uuid.uuid4()))

    def candidate(self):
        book, meta, expected = self.snapshot()
        rows = sheet_values(book["Manual Inputs"])[1:]
        receipts = sheet_values(book["Import Receipts"])[1:]
        received = {(r[1], r[4]) for r in receipts}
        acknowledged_ids = {r[1] for r in receipts}
        entries, _, _ = self.ledger.committed()
        changes, pending, identities = [], [], set()
        alias = next((key for key, root in self.ledger.config["source_roots"].items()
                      if self.path.is_relative_to(__import__("pathlib").Path(root).resolve())), None)
        require(alias is not None, "workbook-source-root-required", "configuration")
        for index, values in enumerate(rows):
            if not any(value is not None for value in values):
                continue
            row = dict(zip(MANUAL, values))
            rid = row["row_id"]
            if rid:
                require(rid not in identities, "duplicate-row-id", "row_id")
                identities.add(rid)
            if (rid, digest(values)) in received:
                continue
            action = row["action"]
            require(action in {"new", "correct", "void"}, "invalid-manual-action", "action")
            require(rid not in acknowledged_ids, "acknowledged-row-edited", "row_id")
            previous = entries.get(row["entry_id"])
            if action == "new":
                require(not row["entry_id"] and not row["base_revision"], "new-row-identity-conflict", "entry")
                entry = {"schema_version": 1, "entry_id": None, "work_id": None,
                         "revision": 1, "previous_revision_hash": None,
                         "owner_workspace": self.ledger.config["workspace_id"], "actor_kind": "human",
                         "visibility": "internal"}
            else:
                require(previous and type(row["base_revision"]) is int and
                        row["base_revision"] == previous["revision"], "stale-base-revision", "entry")
                entry = copy.deepcopy(previous)
                entry.update(revision=previous["revision"] + 1, previous_revision_hash=digest(previous),
                             reason=row["reason"], voided=action == "void")
            if action != "void":
                entry.update({key: row[key] for key in ("performed_on", "client_id", "project_id",
                             "title_internal", "outcome", "status", "activity_kind")})
                entry["outcome"] = row["outcome"] or ""
                entry["timing"] = ({"kind": "interval", "start": row["start"], "end": row["end"],
                                    "timezone": row["timezone"]} if row["start"] or row["end"]
                                   else {"kind": "date", "date": row["performed_on"]})
                entry["duration"] = {"seconds": row["duration_seconds"], "basis": row["basis"],
                                     "evidence_ref": 0 if row["basis"] in ACTUAL_BASES else None,
                                     "session_id": None}
            entry["recorded_at"] = meta["recorded_at"]
            # Stable snapshot timestamp (not clock-now) makes prepare retry deterministic.
            entry["source_refs"] = [{"root_alias": alias,
                "relative_path": self.path.relative_to(__import__("pathlib").Path(self.ledger.config["source_roots"][alias]).resolve()).as_posix(),
                "sha256": expected, "observed_at": entry["recorded_at"],
                "fact_date": entry["performed_on"], "evidence_level": "owner-attested"}]
            changes.append({"row_id": rid, "entry": entry})
            pending.append(index)
        return {"schema_version": 1, "changes": changes,
                "projection_source": {"relative_path": "work-evidence/time-tracker.xlsx", "sha256": expected}}, \
               {"workbook_id": meta["workbook_id"], "input_hash": expected,
                "rows": rows, "receipts": receipts, "pending_indexes": pending}

    def preview(self):
        candidate, _ = self.candidate()
        return self.ledger.preview(candidate) if candidate["changes"] else {"status": "no-new-rows"}

    def prepare(self):
        candidate, projection = self.candidate()
        with self.ledger.lock():
            self.closed()
            data = self.path.read_bytes()
            require(hashlib.sha256(data).hexdigest() == projection["input_hash"] and
                    file_hash(self.path) == projection["input_hash"],
                    "stale-workbook", "workbook")
            atomic_write(self.ledger.runtime("xlsx/backups/" + projection["input_hash"] + ".xlsx"),
                         data, immutable=True)
        result = self.ledger.prepare(candidate)
        with self.ledger.lock():
            atomic_write(self.ledger.runtime("xlsx/pending/" + result["transaction_id"] + ".json"),
                         canonical_bytes(projection), immutable=True)
        return result

    def apply(self, transaction_id, manifest_hash, *, failpoint=None):
        uuid_id(transaction_id, "transaction_id")
        sha(manifest_hash, "manifest_hash")
        self.closed()
        marker = self.ledger.runtime("xlsx/completed/" + transaction_id + ".json")
        if marker.exists():
            require(read_json(marker)["manifest_hash"] == manifest_hash, "stale-approval", "manifest")
            return {"status": "completed", "transaction_id": transaction_id}
        projection = read_json(self.ledger.runtime("xlsx/pending/" + transaction_id + ".json", exists=True))
        manifest = read_json(self.ledger.runtime("transactions/" + transaction_id + "/manifest.json", exists=True))
        require(digest(manifest) == manifest_hash, "stale-approval", "manifest")
        require(manifest["projection_source"]["sha256"] == projection["input_hash"],
                "projection-conflict", "input_hash")
        backup = self.ledger.runtime("xlsx/backups/" + projection["input_hash"] + ".xlsx", exists=True)
        require(file_hash(backup) == projection["input_hash"], "projection-conflict", "backup")
        source_book = load_book(backup.read_bytes())
        original_rows = sheet_values(source_book["Manual Inputs"])[1:]
        original_receipts = sheet_values(source_book["Import Receipts"])[1:]
        received = {(r[1], r[4]) for r in original_receipts}
        indexes = [i for i, row in enumerate(original_rows) if any(v is not None for v in row)
                   and (row[0], digest(row)) not in received]
        source_meta = dict(sheet_values(source_book["Metadata"])[1:])
        require(projection["rows"] == original_rows and projection["receipts"] == original_receipts
                and projection["pending_indexes"] == indexes
                and projection["workbook_id"] == source_meta["workbook_id"]
                and len(indexes) == len(manifest["rows"]),
                "projection-conflict", "pending")
        for index, approved in zip(projection["pending_indexes"], manifest["rows"]):
            projection["rows"][index][0] = approved["row_id"]
            projection["receipts"].append([transaction_id, approved["row_id"],
                approved["entry"]["entry_id"], approved["entry"]["revision"], digest(projection["rows"][index])])
        if file_hash(self.path) != projection["input_hash"]:
            # Replacement succeeded but its receipt write may have crashed. Only
            # the exact protected projection with identical intake can finish.
            book, meta, output_hash = self.snapshot()
            require(meta["workbook_id"] == projection["workbook_id"] and
                    sheet_values(book["Manual Inputs"])[1:] == projection["rows"] and
                    sheet_values(book["Import Receipts"])[1:] == projection["receipts"],
                    "stale-workbook", "workbook")
            _, _, commits = self.ledger.committed()
            require(any(c["transaction_id"] == transaction_id and c["manifest_hash"] == manifest_hash
                        for c in commits), "uncommitted-projection", "workbook")
            with self.ledger.lock():
                atomic_write(marker, canonical_bytes({"manifest_hash": manifest_hash, "output_hash": output_hash}), immutable=True)
            return {"status": "completed", "transaction_id": transaction_id}
        result = self.ledger.apply(transaction_id, manifest_hash)
        if failpoint == "after-canonical":
            raise EvidenceError("injected-failure", "projection")
        with self.ledger.lock():
            view = self.write_view(projection["workbook_id"], projection["rows"], projection["receipts"],
                                   expected=projection["input_hash"])
            if failpoint == "after-workbook":
                raise EvidenceError("injected-failure", "receipt")
            atomic_write(marker, canonical_bytes({"manifest_hash": manifest_hash, "output_hash": view["sha256"]}), immutable=True)
        return {**result, "workbook": view}
